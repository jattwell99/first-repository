"""Ingest ONS mid-year population estimates by local authority.

    python manage.py ingest_population \\
        --path data/gva/populationestimatesbylocalauthority.xlsx \\
        --vintage 2020-06-mye

The vintage identifies the release (e.g. the 2020 mid-year estimates) so
several releases can coexist. As with GVA ingestion, the header row and year
columns are detected heuristically; rows keyed by a valid ONS geography code
are read as whole-number populations.
"""
from __future__ import annotations

from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from economy.models import PopulationValue, PopulationVintage
from geography.models import Geography, GeographyLevel

from ._xlsx import detect_layout, is_geocode, to_number


class Command(BaseCommand):
    help = "Ingest ONS mid-year population estimates by local authority."

    def add_arguments(self, parser):
        parser.add_argument("--path", required=True, help="File or directory.")
        parser.add_argument(
            "--vintage",
            required=True,
            help="Vintage code, e.g. 2020-06-mye.",
        )
        parser.add_argument(
            "--vintage-label",
            default="",
            help="Human label for the vintage (defaults to the code).",
        )
        parser.add_argument(
            "--sheet",
            default="",
            help="Only read sheets whose name contains this text.",
        )
        parser.add_argument(
            "--level",
            default="LAD",
            help="GeographyLevel code for unseen geography codes (default: LAD).",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Parse and report without writing to the database.",
        )

    def handle(self, *args, **options):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise CommandError(
                "openpyxl is required: pip install -r requirements.txt"
            ) from exc

        files = self._resolve_files(options["path"])
        if not files:
            raise CommandError(f"No .xlsx files found at {options['path']}")

        vintage = self._get_vintage(
            options["vintage"], options["vintage_label"], options["dry_run"]
        )
        level = self._get_level(options["level"])

        totals = {"files": 0, "rows": 0, "new_geographies": 0, "years": set()}
        for path in files:
            self.stdout.write(f"Reading {path.name} ...")
            wb = load_workbook(path, read_only=True, data_only=True)
            for ws in wb.worksheets:
                if options["sheet"] and options["sheet"].lower() not in ws.title.lower():
                    continue
                stats = self._ingest_sheet(
                    ws, vintage, level, options["dry_run"]
                )
                if stats["rows"]:
                    totals["rows"] += stats["rows"]
                    totals["new_geographies"] += stats["new_geographies"]
                    totals["years"] |= stats["years"]
                    self.stdout.write(
                        f"  [{ws.title}] {stats['rows']} population values"
                    )
            wb.close()
            totals["files"] += 1

        verb = "Would ingest" if options["dry_run"] else "Ingested"
        years = (
            f"{min(totals['years'])}–{max(totals['years'])}"
            if totals["years"]
            else "no"
        )
        self.stdout.write(
            self.style.SUCCESS(
                f"{verb} {totals['rows']} population values ({years} years) "
                f"for vintage '{options['vintage']}'; "
                f"{totals['new_geographies']} new geographies."
            )
        )

    # -- helpers ----------------------------------------------------------

    def _resolve_files(self, raw_path: str) -> list[Path]:
        path = Path(raw_path)
        if path.is_dir():
            return sorted(path.glob("*.xlsx"))
        if path.is_file():
            return [path]
        return []

    def _get_vintage(self, code, label, dry_run):
        if dry_run:
            return PopulationVintage(code=code, label=label or code)
        vintage, _ = PopulationVintage.objects.get_or_create(
            code=code, defaults={"label": label or code}
        )
        if label and vintage.label != label:
            vintage.label = label
            vintage.save(update_fields=["label"])
        return vintage

    def _get_level(self, code: str) -> GeographyLevel:
        try:
            return GeographyLevel.objects.get(code=code)
        except GeographyLevel.DoesNotExist as exc:
            raise CommandError(
                f"Unknown geography level '{code}'. Run seed_v1 --dimensions first."
            ) from exc

    @transaction.atomic
    def _ingest_sheet(self, ws, vintage, level, dry_run):
        stats = {"rows": 0, "years": set(), "new_geographies": 0}
        layout = detect_layout(ws)
        if layout is None:
            return stats

        for row in ws.iter_rows(min_row=layout.header_row + 1):
            code_cell = row[layout.code_col - 1] if layout.code_col - 1 < len(row) else None
            if code_cell is None or not is_geocode(code_cell.value):
                continue
            code = str(code_cell.value).strip()
            name = ""
            if layout.name_col and layout.name_col - 1 < len(row):
                name = str(row[layout.name_col - 1].value or "").strip()

            geo, created = self._get_or_create_geography(code, name, level, dry_run)
            if created:
                stats["new_geographies"] += 1

            for year, col in layout.year_cols.items():
                if col - 1 >= len(row):
                    continue
                value = to_number(row[col - 1].value)
                if value is None:
                    continue
                stats["years"].add(year)
                stats["rows"] += 1
                if dry_run or geo is None:
                    continue
                PopulationValue.objects.update_or_create(
                    geography=geo,
                    vintage=vintage,
                    year=year,
                    defaults={"population": int(round(value))},
                )
        return stats

    def _get_or_create_geography(self, code, name, level, dry_run):
        geo = Geography.objects.filter(code=code).first()
        if geo is not None:
            return geo, False
        if dry_run:
            return None, True
        geo = Geography.objects.create(code=code, name=name or code, level=level)
        return geo, True
